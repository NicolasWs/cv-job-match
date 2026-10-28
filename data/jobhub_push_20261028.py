import sys; sys.path.insert(0,'/home/nicow/ai-job-search-nico')
from tools.jobhub_sync import job_exists, push_job
import openpyxl
wb = openpyxl.load_workbook('/home/nicow/cv-job-match/applications/Job_Matches_Paris_ProductAI.xlsx')
ws = wb.active
headers = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(headers)}
rows = list(ws.iter_rows(min_row=2, values_only=True))
pushed = 0
skipped = 0
errors = 0
for r in rows:
    try:
        fsv = float(r[idx['Fit Score (/10)']])
    except Exception:
        continue
    if fsv < 7.5:
        continue
    company = r[idx['Company']] or ''
    title = r[idx['Job Title']] or ''
    if not company or not title:
        continue
    try:
        ex = job_exists(company, title)
        if ex.get('exists'):
            skipped += 1
            continue
        res = push_job(
            company=company, title=title, location=r[idx['Location']] or '',
            source_url=r[idx['Job Posting URL']] or '', source='cv-coach-scan',
            fit_pct=round(fsv * 10),
            match_rationale=(r[idx["Why this fits (or doesn't)"]] or '')[:300],
        )
        if res is None:
            errors += 1
        else:
            pushed += 1
            print('pushed', company, title, fsv)
    except Exception as e:
        errors += 1
        print('ERR', company, title, e)
print('done pushed=', pushed, 'skipped=', skipped, 'errors=', errors)
